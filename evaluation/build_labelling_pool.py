"""Build the pooled, blind relevance-labelling sheet (benchmark v2).

Run from the project root, AFTER every method's Top-5 file exists:
    python evaluation/build_labelling_pool.py
    python evaluation/build_labelling_pool.py --no-ablation     # methods only

Pool
    For every validation/test query whose eligibility role is "ranking" (E >= 5), the union of
    the Top-5 lists of all evaluated methods:
        cpg (outputs/top5/cpg_top5_recommendations.csv)
        content_based, rule_based (outputs/baseline/*_top5.csv)
        lightgcn seeds 42/43/44 (baseline_models/lightgcn_outputs/lightgcn_top5_recommendations.csv)
        ablation / robustness conditions (outputs/ablation/top5_all_conditions.csv, optional)
    Each (query, property) pair appears ONCE, whichever and however many methods returned it.

Blinding
    The labelling sheet shows only the search profile and property information. Method names,
    ranks, scores and how many methods returned a property are NOT shown; properties are
    shuffled within each query (seed 2026). The pair -> method mapping is written to a separate
    key file that the assessor must not open.

Outputs (evaluation/labelling_v2/):
    relevance_labelling_sheet.xlsx      for the primary assessor (all pairs)
    relevance_labelling_sheet_A2.xlsx   for the second assessor (subset of queries, same layout)
    pool_key_DO_NOT_SHOW_ASSESSOR.csv   pair_id -> query, Property_ID, methods that returned it
    pool_summary.json                   counts, sources, subset queries, seed
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import recommender_service as svc  # noqa: E402

QUERIES = ROOT / "frozen_queries" / "frozen_queries_set_v2.csv"
MANIFEST = ROOT / "benchmark" / "eligibility_manifest_v2.csv"
SOURCES = {
    "cpg": ROOT / "outputs" / "top5" / "cpg_top5_recommendations.csv",
    "content_based": ROOT / "outputs" / "baseline" / "content_based_top5.csv",
    "rule_based": ROOT / "outputs" / "baseline" / "rule_based_top5.csv",
    "lightgcn": ROOT / "baseline_models" / "lightgcn_outputs" / "lightgcn_top5_recommendations.csv",
}
ABLATION = ROOT / "outputs" / "ablation" / "top5_all_conditions.csv"
OUT_DIR = ROOT / "evaluation" / "labelling_v2"
SEED = 2026
SECOND_ASSESSOR_SHARE = 0.25
EVAL_PARTITIONS = ("validation", "test")

FACILITIES = [("train_station_count_3km", "Train stations ≤3 km"), ("school_count_3km", "Schools ≤3 km"),
              ("shopping_mall_count_3km", "Malls ≤3 km"), ("supermarket_count_3km", "Supermarkets ≤3 km"),
              ("park_count_1km", "Parks ≤1 km")]

RUBRIC = [
    ("3", "Highly relevant",
     "Fits the search profile well overall: strong on the stated main priority and on the soft "
     "preferences (preferred size, selected facilities), with nothing that would make a buyer with "
     "this profile hesitate."),
    ("2", "Relevant",
     "A reasonable option for this profile: fits the main priority and most preferences, with minor "
     "compromises a typical buyer would accept."),
    ("1", "Marginal",
     "Meets the requirements but is a weak fit: clearly misses the main priority or several "
     "preferences; a buyer would consider it only if few alternatives exist."),
    ("0", "Not relevant",
     "Would not be a sensible suggestion for this profile despite meeting the hard requirements "
     "(e.g. size far from preferred, none of the selected facilities nearby, price at the limit "
     "when price is the main priority)."),
]
PROTOCOL_ID = "REL-v2-1"
PROTOCOL = [
    ("Protocol", f"{PROTOCOL_ID}, frozen 2026-10-31, before any label was entered."),
    ("Unit of assessment", "One listing for one search profile (a query-property pair). Near-identical "
                           "listings are judged individually."),
    ("Relevant (binary)", "Label 2 or 3. Used for Precision@5, Recall@5 and MRR; NDCG@5 uses the 0-3 grades."),
    ("Assessors", "A1 (researcher) labels every pair. A2 (second assessor) labels a random 25% of the "
                  "queries independently, without seeing A1's labels or any recommender output."),
    ("Agreement", "Linear-weighted Cohen's kappa on the double-labelled pairs, computed on the original "
                  "labels before adjudication; also reported: exact agreement % and unweighted kappa on "
                  "the binary label. Target kappa >= 0.60."),
    ("Disagreements", "Every pair where A1 and A2 differ is resolved by discussion to a consensus label, "
                      "using only this sheet (no recommender output). If no consensus is reached, the "
                      "lower of the two labels is used. Both original labels are kept on record."),
    ("After freezing", "Labels are not changed after evaluation results are seen. If kappa is below target "
                       "it is reported as a limitation."),
]


def load_sources(include_ablation: bool) -> pd.DataFrame:
    frames = []
    for name, path in SOURCES.items():
        if not path.exists():
            raise FileNotFoundError(f"Missing Top-5 file for {name}: {path}")
        d = pd.read_csv(path).dropna(subset=["Property_ID"])
        if name == "lightgcn":
            d["source"] = "lightgcn_seed" + d["seed"].astype(int).astype(str)
        else:
            d["source"] = name
        frames.append(d[["query_id", "Property_ID", "source"]])
    if include_ablation and ABLATION.exists():
        a = pd.read_csv(ABLATION).dropna(subset=["Property_ID"])
        a["source"] = "ablation:" + a["condition"].astype(str)
        frames.append(a[["query_id", "Property_ID", "source"]])
    out = pd.concat(frames, ignore_index=True)
    out["Property_ID"] = out["Property_ID"].astype(int)
    return out


def profile_text(q: pd.Series) -> dict:
    def yes(v):
        return str(v).strip().lower() in {"true", "1", "yes"}
    facilities = svc._facilities(q.get("nearby_facilities")) or ["no preference"]
    return {
        "Area (district)": str(q["preferred_area"]).title(),
        "Max budget": f"RM {float(q['max_budget_rm']):,.0f}",
        "Property type": q["property_type"] or "Any",
        "Min bedrooms": q["min_bedrooms"] or "–",
        "Min bathrooms": q["min_bathrooms"] or "–",
        "Parking required": "Yes" if yes(q["parking_required"]) else "No",
        "Furnished only": "Yes" if yes(q["furnished_only"]) else "No",
        "Tenure": q["tenure"] or "No preference",
        "Main priority": str(q["main_priority"]).title(),
        "Nearby facilities wanted": ", ".join(facilities),
        "Preferred size": f"{float(q['preferred_built_up_sqft']):,.0f} sq ft"
        if str(q["preferred_built_up_sqft"]).strip() not in {"", "0", "0.0"} else "–",
    }


HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
QUERY_FILL = PatternFill("solid", fgColor="E8EEF6")
LABEL_FILL = PatternFill("solid", fgColor="FFF4CC")
WRAP = Alignment(wrap_text=True, vertical="top")

COLUMNS = [("pair_id", 11), ("Property", 30), ("Area", 22), ("Type", 16), ("Price (RM)", 12),
           ("Built-up (sq ft)", 11), ("Bed", 6), ("Bath", 6), ("Car parks", 8), ("Furnishing", 15),
           ("Tenure", 11), *[(lbl, 10) for _, lbl in FACILITIES], ("Listing link", 14),
           ("Relevance (0-3)", 11), ("Comment (optional)", 30)]


def write_instructions(ins, assessor: str) -> None:
    """Instructions tab: task, frozen 0-3 rubric and annotation protocol."""
    lines = [
        (f"Relevance labelling – benchmark v2 ({PROTOCOL_ID})", True),
        (f"Assessor: {assessor}", False),
        ("", False),
        ("For each search profile, judge how well each listed property suits a buyer with that profile.", False),
        ("Every property shown already meets all hard requirements (area, budget, type, rooms, parking, "
         "furnishing, tenure). Judge the overall fit, not whether it passes the requirements.", False),
        ("Judge each property on its own. The order of properties is random and says nothing about quality. "
         "Near-identical listings are judged individually.", False),
        ("Use only the information in this sheet (you may open the listing link). Do not look at any "
         "recommender output, score or ranking while labelling.", False),
        ("Enter one label per row in the yellow column. Leave nothing blank.", False),
        ("", False),
        ("Relevance scale (0-3)", True),
    ]

    def full_width(text, font):
        ins.append([text])
        r = ins.max_row
        ins.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ins.cell(r, 1).font = font
        ins.cell(r, 1).alignment = WRAP
        if len(text) > 120:
            ins.row_dimensions[r].height = 32

    def pair_row(a, b):
        ins.append([a, b])
        r = ins.max_row
        ins.cell(r, 1).font = Font(bold=True)
        ins.cell(r, 1).alignment = WRAP
        ins.cell(r, 2).alignment = WRAP
        ins.row_dimensions[r].height = 32

    for text, bold in lines:
        full_width(text, Font(bold=bold, size=13 if bold else 11))
    for lv, name, desc in RUBRIC:
        pair_row(f"{lv} – {name}", desc)
    ins.append([])
    full_width("Annotation protocol", Font(bold=True, size=13))
    for a, b in PROTOCOL:
        pair_row(a, b)
    ins.column_dimensions["A"].width = 22
    ins.column_dimensions["B"].width = 110


def refresh_instructions(path: Path, assessor: str) -> None:
    """Replace only the Instructions tab of an existing sheet (labels are kept)."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    if "Instructions" in wb.sheetnames:
        del wb["Instructions"]
    ins = wb.create_sheet("Instructions", 0)
    write_instructions(ins, assessor)
    wb.active = 0
    wb.save(path)


def write_sheet(path: Path, pool: pd.DataFrame, queries: pd.DataFrame, catalogue: pd.DataFrame,
                assessor: str) -> None:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    write_instructions(ins, assessor)

    ws = wb.create_sheet("Labels")
    ws.append([c for c, _ in COLUMNS])
    for i, (_, w) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    label_col = [c for c, _ in COLUMNS].index("Relevance (0-3)") + 1
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="3", allow_blank=True,
                        showErrorMessage=True, errorTitle="Relevance", error="Enter 0, 1, 2 or 3.")
    ws.add_data_validation(dv)

    info = catalogue.set_index("Property_ID")
    qinfo = queries.set_index("query_id")
    for qid, grp in pool.groupby("query_id", sort=True):
        prof = profile_text(qinfo.loc[qid])
        text = f"SEARCH PROFILE {qid}   " + "   |   ".join(f"{k}: {v}" for k, v in prof.items())
        ws.append([text])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLUMNS))
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 1).fill = QUERY_FILL
        ws.cell(r, 1).alignment = WRAP
        ws.row_dimensions[r].height = 48
        for _, p in grp.iterrows():
            c = info.loc[p["Property_ID"]]
            link = str(c.get("original_link") or "")
            ws.append([p["pair_id"], svc._name(c), f"{str(c['area']).title()}, {str(c['district']).title()}",
                       c["property_type"], float(c["price"]), float(c["built_up_area"]),
                       int(c["bedroom"]), int(c["bathroom"]),
                       None if pd.isna(c["car_parks"]) else int(c["car_parks"]), c["furnishing"], c["tenure"],
                       *[None if pd.isna(c[col]) else int(c[col]) for col, _ in FACILITIES],
                       "open" if link else "", None, None])
            r = ws.max_row
            ws.cell(r, 5).number_format = "#,##0"
            ws.cell(r, 6).number_format = "#,##0"
            if link:
                ws.cell(r, len(COLUMNS) - 2).hyperlink = link
                ws.cell(r, len(COLUMNS) - 2).style = "Hyperlink"
            ws.cell(r, label_col).fill = LABEL_FILL
            dv.add(ws.cell(r, label_col))
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(r, col).alignment = WRAP
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description="Pooled blind relevance-labelling sheet (v2).")
    ap.add_argument("--no-ablation", action="store_true", help="Pool the methods only.")
    ap.add_argument("--force", action="store_true", help="Overwrite an existing labelling folder.")
    ap.add_argument("--refresh-instructions", action="store_true",
                    help="Only rewrite the Instructions tab of the existing sheets (labels are kept).")
    args = ap.parse_args()
    if args.refresh_instructions:
        for name, who in (("relevance_labelling_sheet.xlsx", "primary (A1)"),
                          ("relevance_labelling_sheet_A2.xlsx", "second assessor (A2) – agreement subset")):
            refresh_instructions(OUT_DIR / name, who)
        print(f"Instructions updated to {PROTOCOL_ID} in {OUT_DIR}")
        return
    if (OUT_DIR / "relevance_labelling_sheet.xlsx").exists() and not args.force:
        raise SystemExit(f"{OUT_DIR} already has a labelling sheet. Labels may already be entered; "
                         "use --force only if you really want to rebuild it.")

    queries = pd.read_csv(QUERIES, dtype=str, keep_default_na=False)
    manifest = pd.read_csv(MANIFEST, dtype=str, keep_default_na=False)
    eval_q = manifest[(manifest["benchmark_role"] == "ranking")
                      & manifest["partition"].isin(EVAL_PARTITIONS)]["query_id"]
    catalogue = pd.read_csv(svc.DATA_PATH)

    src = load_sources(not args.no_ablation)
    src = src[src["query_id"].isin(set(eval_q))]
    missing = set(src["Property_ID"]) - set(catalogue["Property_ID"])
    if missing:
        raise ValueError(f"{len(missing)} pooled properties are not in the catalogue, e.g. {sorted(missing)[:5]}")

    key = (src.groupby(["query_id", "Property_ID"])["source"]
           .apply(lambda s: ";".join(sorted(set(s)))).reset_index(name="returned_by"))
    key["n_sources"] = key["returned_by"].str.count(";") + 1
    key = key.merge(manifest[["query_id", "partition", "eligible_count"]], on="query_id")

    # Shuffle within each query, then give anonymous ids in sheet order.
    rng = np.random.default_rng(SEED)
    key["_r"] = rng.random(len(key))
    key = key.sort_values(["query_id", "_r"]).drop(columns="_r").reset_index(drop=True)
    key.insert(0, "pair_id", [f"P{i:04d}" for i in range(1, len(key) + 1)])

    # Second-assessor subset: whole queries, stratified by partition.
    sub = []
    for part, g in key.drop_duplicates("query_id").groupby("partition"):
        n = max(1, round(len(g) * SECOND_ASSESSOR_SHARE))
        sub += sorted(rng.choice(g["query_id"].to_numpy(), size=n, replace=False).tolist())
    key["second_assessor"] = key["query_id"].isin(sub)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    key.to_csv(OUT_DIR / "pool_key_DO_NOT_SHOW_ASSESSOR.csv", index=False)
    write_sheet(OUT_DIR / "relevance_labelling_sheet.xlsx", key, queries, catalogue, "primary (A1)")
    write_sheet(OUT_DIR / "relevance_labelling_sheet_A2.xlsx", key[key["second_assessor"]], queries,
                catalogue, "second assessor (A2) – agreement subset")

    per_source = src.drop_duplicates(["query_id", "Property_ID", "source"])["source"].value_counts()
    summary = {
        "benchmark": "v2", "protocol": PROTOCOL_ID, "seed": SEED, "partitions": EVAL_PARTITIONS,
        "queries": int(key["query_id"].nunique()),
        "queries_by_partition": key.drop_duplicates("query_id")["partition"].value_counts().to_dict(),
        "pairs_to_label": int(len(key)),
        "pairs_per_query": {"mean": round(float(key.groupby("query_id").size().mean()), 2),
                            "min": int(key.groupby("query_id").size().min()),
                            "max": int(key.groupby("query_id").size().max())},
        "pairs_from_methods_only": int(key["returned_by"].str.split(";").apply(
            lambda s: any(not x.startswith("ablation:") for x in s)).sum()),
        "pairs_added_by_ablation_only": int(key["returned_by"].str.split(";").apply(
            lambda s: all(x.startswith("ablation:") for x in s)).sum()),
        "pairs_returned_per_source": {k: int(v) for k, v in per_source.items() if not k.startswith("ablation:")},
        "second_assessor": {"share_of_queries": SECOND_ASSESSOR_SHARE, "queries": sub,
                            "pairs": int(key["second_assessor"].sum())},
    }
    (OUT_DIR / "pool_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))
    print(f"\nSaved to {OUT_DIR}")


if __name__ == "__main__":
    main()
